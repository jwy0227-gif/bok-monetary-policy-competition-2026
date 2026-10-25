from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "데이터 정리.xlsx"
OUT = ROOT / "data" / "raw" / "data_workbook_features.json"


def cell_value(value: object) -> object:
    if value is None:
        return None
    return value


def main() -> int:
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
    sheets: list[dict[str, object]] = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        feature_rows: list[dict[str, object]] = []
        header_row_idx = None
        headers: list[object] = []
        for idx, row in enumerate(rows, start=1):
            first = row[0] if row else None
            second = row[1] if len(row) > 1 else None
            if first == "지표" and second == "단위":
                header_row_idx = idx
                headers = list(row)
                break
        if header_row_idx is not None:
            for row_idx, row in enumerate(rows[header_row_idx:], start=header_row_idx + 1):
                name = row[0] if row else None
                unit = row[1] if len(row) > 1 else None
                if not name or str(name).strip() in {"", "nan"}:
                    continue
                values = {}
                for col_idx, header in enumerate(headers[2:], start=2):
                    if header is None:
                        continue
                    value = row[col_idx] if col_idx < len(row) else None
                    values[str(header)] = cell_value(value)
                feature_rows.append(
                    {
                        "row": row_idx,
                        "name": str(name).strip(),
                        "unit": None if unit is None else str(unit).strip(),
                        "values": values,
                    }
                )
        preview = []
        for row in rows[:8]:
            preview.append([cell_value(v) for v in row[:8]])
        sheets.append(
            {
                "sheet": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "header_row": header_row_idx,
                "features": feature_rows,
                "preview": preview,
            }
        )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(sheets, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(OUT)
    for sheet in sheets:
        print(f"{sheet['sheet']}: {len(sheet['features'])} features")
        for feature in sheet["features"][:3]:
            print(f"  - {feature['name']} ({feature['unit']})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
