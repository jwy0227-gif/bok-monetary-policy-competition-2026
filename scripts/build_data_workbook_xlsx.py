from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
FEATURES = ROOT / "data" / "raw" / "data_workbook_features.json"
PANEL = ROOT / "data" / "processed" / "data_workbook_bok_panel_2010.csv"
STATUS = ROOT / "data" / "processed" / "data_workbook_collection_status.csv"
UNRESOLVED = ROOT / "data" / "processed" / "data_workbook_unresolved_features_with_sources.csv"
OUT_DIR = ROOT / "outputs" / "data_workbook_organized"
OUT_XLSX = OUT_DIR / "데이터_정리_2010_수집본.xlsx"
FALLBACK_XLSX = OUT_DIR / "데이터_정리_2010_수집본_보강.xlsx"


STATUS_LABELS = {
    "BOK_OK": "BOK 수집",
    "CALCULATED_OK": "계산",
    "EXTERNAL_OK": "외부 수집",
    "NOT_COLLECTED": "미수집",
    "BOK_ERROR": "BOK 오류",
    "NO_BOK_ROWS": "BOK 빈값",
    "CALC_NO_ROWS": "계산 빈값",
}

STATUS_FILLS = {
    "BOK_OK": "E2F0D9",
    "CALCULATED_OK": "D9EAF7",
    "EXTERNAL_OK": "EADCF8",
    "NOT_COLLECTED": "FCE4D6",
    "BOK_ERROR": "F4CCCC",
    "NO_BOK_ROWS": "F4CCCC",
    "CALC_NO_ROWS": "F4CCCC",
}

CATEGORY_NOTES = {
    "02.기준금리·통화": "기준금리, 정책금리, 통화량, 유동성 및 금리차 계산 지표",
    "03.물가": "CPI/PPI, 기대인플레이션, 품목별 물가 및 물가 관련 지표",
    "04.경기·성장": "GDP/GNI, 산업활동, 심리지수 및 기업경기 지표",
    "05.고용": "고용, 실업, 청년고용, 임금 관련 지표",
    "06.국제수지·교역": "국제수지와 수출입/품목별 교역 관련 지표",
    "07.금융·외환시장": "채권금리, 신용스프레드, 환율, 주식시장 및 외국인 자금흐름",
    "08.가계·기업신용": "가계신용, 주택관련대출, 기업대출 및 연체율",
    "09.부동산·금융안정": "주택가격, 금융안정, 외환보유액, 외채, PF 관련 지표",
    "10.세계경제": "세계/주요국 성장률, 해외 정책금리, 원유 및 글로벌 가격지표",
}

EXTRA_FEATURES = {
    "06.국제수지·교역": [
        ("수출금액지수(전년동월비)", "%"),
        ("수출물량지수(전년동월비)", "%"),
        ("수입금액지수(전년동월비)", "%"),
        ("수입물량지수(전년동월비)", "%"),
    ],
    "07.금융·외환시장": [
        ("채권거래 합계(거래대금)", "조원"),
        ("회사채 거래대금", "조원"),
    ],
}


def safe_sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\[\]\*:/\\?]", "·", name)
    return cleaned[:31]


def load_features() -> list[dict]:
    sheets = json.loads(FEATURES.read_text(encoding="utf-8"))
    return [sheet for sheet in sheets if sheet.get("features")]


def normalize_status(status: pd.DataFrame) -> dict[str, dict[str, str]]:
    priority = {
        "BOK_OK": 1,
        "CALCULATED_OK": 2,
        "EXTERNAL_OK": 3,
        "NOT_COLLECTED": 9,
        "NO_BOK_ROWS": 8,
        "CALC_NO_ROWS": 8,
        "BOK_ERROR": 8,
    }
    status = status.copy()
    status["_priority"] = status["status"].map(priority).fillna(99)
    status = status.sort_values(["feature", "_priority"]).drop_duplicates("feature", keep="first")
    return {
        row["feature"]: {
            "status": row.get("status", ""),
            "variable": "" if pd.isna(row.get("variable")) else str(row.get("variable")),
            "source": "" if pd.isna(row.get("source")) else str(row.get("source")),
            "detail": "" if pd.isna(row.get("detail")) else str(row.get("detail")),
        }
        for _, row in status.iterrows()
    }


def build_value_maps(panel: pd.DataFrame) -> tuple[list[str], dict[str, dict[str, float]], dict[str, tuple[str, float]]]:
    panel = panel.copy()
    panel["date"] = pd.to_datetime(panel["date"])
    panel["month"] = panel["date"].dt.strftime("%Y-%m")
    month_labels = pd.date_range(panel["date"].min(), panel["date"].max(), freq="MS").strftime("%Y-%m").tolist()
    values: dict[str, dict[str, float]] = {}
    latest: dict[str, tuple[str, float]] = {}
    for feature, group in panel.sort_values(["date"]).groupby("feature"):
        month_map = group.drop_duplicates("month", keep="last").set_index("month")["value"].to_dict()
        values[feature] = month_map
        last = group.dropna(subset=["value"]).tail(1)
        if not last.empty:
            latest[feature] = (last.iloc[0]["month"], float(last.iloc[0]["value"]))
    return month_labels, values, latest


def style_sheet(ws, max_row: int, max_col: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    title_fill = PatternFill("solid", fgColor="17365D")
    white = "FFFFFF"
    border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )
    ws["A1"].fill = title_fill
    ws["A1"].font = Font(bold=True, color=white, size=14)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24
    for cell in ws[4]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=5, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, float):
                cell.number_format = "0.00"
    ws.freeze_panes = "G5"
    ws.auto_filter.ref = f"A4:{get_column_letter(max_col)}{max_row}"


def set_widths(ws, max_col: int) -> None:
    widths = {
        "A": 30,
        "B": 10,
        "C": 12,
        "D": 24,
        "E": 26,
        "F": 42,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for col_idx in range(7, max_col + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 10


def add_category_sheet(
    wb: Workbook,
    sheet: dict,
    month_labels: list[str],
    value_map: dict[str, dict[str, float]],
    latest: dict[str, tuple[str, float]],
    status_map: dict[str, dict[str, str]],
) -> None:
    ws = wb.create_sheet(safe_sheet_name(sheet["sheet"]))
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws["A1"] = sheet["sheet"]
    ws["A2"] = CATEGORY_NOTES.get(sheet["sheet"], "")
    ws["A2"].font = Font(italic=True, color="666666")
    headers = ["지표", "단위", "수집상태", "변수명", "출처", "비고", *month_labels, "마지막 기간", "마지막 값"]
    ws.append([])
    ws.append(headers)

    for feature in sheet["features"]:
        name = feature["name"]
        if name.startswith("※"):
            continue
        meta = status_map.get(name, {"status": "NOT_COLLECTED", "variable": "", "source": "", "detail": ""})
        status = meta["status"]
        row = [
            name,
            feature.get("unit") or "",
            STATUS_LABELS.get(status, status),
            meta.get("variable", ""),
            meta.get("source", ""),
            meta.get("detail", ""),
        ]
        values = value_map.get(name, {})
        row.extend(values.get(month) for month in month_labels)
        last_month, last_value = latest.get(name, ("", None))
        row.extend([last_month, last_value])
        ws.append(row)
        status_cell = ws.cell(row=ws.max_row, column=3)
        status_cell.fill = PatternFill("solid", fgColor=STATUS_FILLS.get(status, "FFFFFF"))
        status_cell.alignment = Alignment(horizontal="center")

    for name, unit in EXTRA_FEATURES.get(sheet["sheet"], []):
        if name not in value_map:
            continue
        meta = status_map.get(name, {"status": "BOK_OK", "variable": "", "source": "", "detail": ""})
        status = meta["status"]
        row = [
            name,
            unit,
            STATUS_LABELS.get(status, status),
            meta.get("variable", ""),
            meta.get("source", ""),
            f"BOK 추가수집 · {meta.get('detail', '')}",
        ]
        values = value_map.get(name, {})
        row.extend(values.get(month) for month in month_labels)
        last_month, last_value = latest.get(name, ("", None))
        row.extend([last_month, last_value])
        ws.append(row)
        ws.cell(row=ws.max_row, column=1).fill = PatternFill("solid", fgColor="FFF2CC")
        status_cell = ws.cell(row=ws.max_row, column=3)
        status_cell.fill = PatternFill("solid", fgColor=STATUS_FILLS.get(status, "FFFFFF"))
        status_cell.alignment = Alignment(horizontal="center")

    max_row = ws.max_row
    max_col = ws.max_column
    style_sheet(ws, max_row, max_col)
    set_widths(ws, max_col)


def add_intro_sheet(wb: Workbook, panel: pd.DataFrame, status: pd.DataFrame, month_labels: list[str]) -> None:
    ws = wb.active
    ws.title = "00.안내"
    ws["A1"] = "데이터 정리 2010년 이후 수집본"
    ws["A1"].font = Font(bold=True, size=16, color="17365D")
    lines = [
        ("생성 기준", "data/processed/data_workbook_bok_panel_2010.csv"),
        ("기간", f"{month_labels[0]} ~ {month_labels[-1]}"),
        ("패널 행 수", len(panel)),
        ("수집 피처 수", panel["feature"].nunique()),
        ("형식", "원본 데이터 정리.xlsx의 지표/단위/기간/출처 구조를 확장"),
        ("주의", "분기 지표는 해당 분기의 시작월 열에 값이 표시됩니다."),
    ]
    ws.append([])
    for key, value in lines:
        ws.append([key, value])
    ws.append([])
    ws.append(["수집상태", "개수"])
    for state, count in status["status"].value_counts().items():
        ws.append([STATUS_LABELS.get(state, state), int(count)])
    for col in ["A", "B"]:
        ws.column_dimensions[col].width = 34
    for cell in ws["A"]:
        cell.font = Font(bold=True)


def add_status_sheet(wb: Workbook, status: pd.DataFrame) -> None:
    ws = wb.create_sheet("01.수집현황")
    ws["A1"] = "피처별 수집현황"
    ws["A1"].font = Font(bold=True, size=14, color="17365D")
    headers = ["지표", "상태", "변수명", "출처", "상세"]
    ws.append([])
    ws.append([])
    ws.append(headers)
    for _, row in status.sort_values(["status", "feature"]).iterrows():
        if str(row["feature"]).startswith("_"):
            continue
        state = row["status"]
        ws.append(
            [
                row["feature"],
                STATUS_LABELS.get(state, state),
                "" if pd.isna(row.get("variable")) else row.get("variable"),
                "" if pd.isna(row.get("source")) else row.get("source"),
                "" if pd.isna(row.get("detail")) else row.get("detail"),
            ]
        )
        ws.cell(row=ws.max_row, column=2).fill = PatternFill("solid", fgColor=STATUS_FILLS.get(state, "FFFFFF"))
    for idx, width in enumerate([34, 14, 28, 32, 80], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    style_sheet(ws, ws.max_row, ws.max_column)
    ws.freeze_panes = "A4"


def add_unresolved_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("11.미수집·외부필요")
    ws["A1"] = "BOK/외부 자동수집 미완료 지표"
    ws["A1"].font = Font(bold=True, size=14, color="17365D")
    unresolved = pd.read_csv(UNRESOLVED)
    headers = ["시트", "지표", "단위", "상태", "원 출처", "원 비고", "상세"]
    ws.append([])
    ws.append([])
    ws.append(headers)
    for _, row in unresolved.iterrows():
        ws.append([row.get(col, "") for col in ["sheet", "feature", "unit", "status", "workbook_source", "workbook_note", "detail"]])
    for idx, width in enumerate([20, 34, 10, 16, 22, 48, 56], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    style_sheet(ws, ws.max_row, ws.max_column)
    ws.freeze_panes = "A4"


def add_tables(wb: Workbook) -> None:
    for ws in wb.worksheets:
        if ws.max_row < 4 or ws.max_column < 2:
            continue
        ref = f"A4:{get_column_letter(ws.max_column)}{ws.max_row}"
        table_name = re.sub(r"[^A-Za-z0-9_]", "_", f"T_{ws.title}")[:240]
        if not re.match(r"^[A-Za-z_]", table_name):
            table_name = f"T_{table_name}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        try:
            ws.add_table(table)
        except ValueError:
            pass


def main() -> int:
    sheets = load_features()
    panel = pd.read_csv(PANEL)
    status = pd.read_csv(STATUS)
    month_labels, value_map, latest = build_value_maps(panel)
    status_map = normalize_status(status)

    wb = Workbook()
    add_intro_sheet(wb, panel, status[~status["feature"].astype(str).str.startswith("_")], month_labels)
    add_status_sheet(wb, status)
    for sheet in sheets:
        if sheet["sheet"] in {"11.출처·메모"}:
            continue
        add_category_sheet(wb, sheet, month_labels, value_map, latest, status_map)
    add_unresolved_sheet(wb)
    add_tables(wb)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUT_XLSX
    try:
        wb.save(output_path)
    except PermissionError:
        output_path = FALLBACK_XLSX
        wb.save(output_path)

    # Verification pass: ensure Excel can reopen and expected sheets exist.
    check = load_workbook(output_path, read_only=True, data_only=True)
    print(output_path)
    print(f"sheets={len(check.sheetnames)} rows_panel={len(panel)} features={panel['feature'].nunique()}")
    print(", ".join(check.sheetnames))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
